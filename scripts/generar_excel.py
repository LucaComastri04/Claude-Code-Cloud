"""Convierte un JSON de prospectos en un Excel formateado y actualiza el histórico.

Uso: python3 scripts/generar_excel.py salida/prospectos_AAAA-MM-DD.json
"""

import csv
import json
import re
import sys
import unicodedata
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

RAIZ = Path(__file__).resolve().parent.parent
HISTORICO = RAIZ / "salida" / "historico.csv"

COLUMNAS = [
    # (cabecera, clave, ancho)
    ("Prioridad", "prioridad", 10),
    ("Puntuación", "puntuacion", 11),
    ("Empresa", "empresa", 28),
    ("Web", "web", 28),
    ("Sector", "sector", 14),
    ("Subsector", "subsector", 18),
    ("País sede", "pais_sede", 12),
    ("Países operación", "paises_operacion", 22),
    ("Empleados", "empleados", 18),
    ("Facturación", "facturacion", 16),
    ("Señales de compra", "senales", 55),
    ("Decisor", "decisor_nombre", 22),
    ("Cargo decisor", "decisor_cargo", 22),
    ("URL decisor", "decisor_url", 30),
    ("Punto de entrada", "entrada_nombre", 22),
    ("Cargo entrada", "entrada_cargo", 22),
    ("URL entrada", "entrada_url", 30),
    ("Por qué encaja", "por_que_encaja", 55),
    ("Servicio SerBan", "servicio_serban", 26),
    ("Ángulo de apertura", "angulo_apertura", 45),
    ("Parecido a", "parecido_a", 30),
    ("Desglose puntuación", "desglose", 34),
    ("Notas", "notas", 35),
    ("Estado", None, 26),
]

COLOR_PRIORIDAD = {"A": "F8CBAD", "B": "FFE699", "C": "FFF2CC"}
FORMAS_JURIDICAS = r"\b(s\.?a\.?u?|s\.?l\.?u?|s\.?a\.? de c\.?v\.?|sas|ltda|inc|llc|grupo|group|sociedad anonima)\b"


def normalizar(nombre: str) -> str:
    texto = unicodedata.normalize("NFKD", nombre.lower())
    texto = "".join(c for c in texto if not unicodedata.combining(c))
    texto = re.sub(FORMAS_JURIDICAS, " ", texto)
    return re.sub(r"[^a-z0-9]+", " ", texto).strip()


def formatear_senales(senales) -> str:
    if isinstance(senales, str):
        return senales
    lineas = []
    for s in senales or []:
        lineas.append(f"• {s.get('senal', '')} ({s.get('fecha', '')}) {s.get('fuente', '')}".strip())
    return "\n".join(lineas)


def cargar_historico() -> set:
    if not HISTORICO.exists():
        return set()
    with HISTORICO.open(encoding="utf-8") as f:
        return {normalizar(fila["empresa"]) for fila in csv.DictReader(f)}


def main() -> None:
    if len(sys.argv) != 2:
        sys.exit(__doc__)
    origen = Path(sys.argv[1])
    prospectos = json.loads(origen.read_text(encoding="utf-8"))

    hay_crm = any(f.name != "LEEME.md" for f in (RAIZ / "datos" / "crm").glob("*"))
    estado_inicial = "Pendiente" if hay_crm else "Pendiente – comprobar en CRM"

    ya_vistas = cargar_historico()
    nuevos, repetidos = [], []
    for p in prospectos:
        (repetidos if normalizar(p.get("empresa", "")) in ya_vistas else nuevos).append(p)
    nuevos.sort(key=lambda p: -int(p.get("puntuacion") or 0))
    if not nuevos:
        sys.exit("Todas las empresas ya están en el histórico; no se genera Excel.")

    wb = Workbook()
    ws = wb.active
    ws.title = "Prospectos"
    cabecera_fill = PatternFill("solid", fgColor="1F3864")
    for col, (titulo, _, ancho) in enumerate(COLUMNAS, start=1):
        celda = ws.cell(row=1, column=col, value=titulo)
        celda.font = Font(bold=True, color="FFFFFF")
        celda.fill = cabecera_fill
        celda.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(col)].width = ancho

    for fila, p in enumerate(nuevos, start=2):
        for col, (_, clave, _) in enumerate(COLUMNAS, start=1):
            if clave is None:
                valor = estado_inicial
            elif clave == "senales":
                valor = formatear_senales(p.get(clave))
            else:
                valor = p.get(clave, "")
            celda = ws.cell(row=fila, column=col, value=valor)
            celda.alignment = Alignment(vertical="top", wrap_text=True)
            if isinstance(valor, str) and valor.startswith("http"):
                celda.hyperlink = valor
                celda.font = Font(color="0563C1", underline="single")
        color = COLOR_PRIORIDAD.get(str(p.get("prioridad", "")).upper())
        if color:
            ws.cell(row=fila, column=1).fill = PatternFill("solid", fgColor=color)

    ws.freeze_panes = "D2"
    ws.auto_filter.ref = ws.dimensions

    destino = origen.with_suffix(".xlsx")
    wb.save(destino)

    hoy = date.today().isoformat()
    nuevo_archivo = not HISTORICO.exists()
    with HISTORICO.open("a", encoding="utf-8", newline="") as f:
        escritor = csv.writer(f)
        if nuevo_archivo:
            escritor.writerow(["fecha", "empresa", "prioridad", "puntuacion"])
        for p in nuevos:
            escritor.writerow([hoy, p.get("empresa", ""), p.get("prioridad", ""), p.get("puntuacion", "")])

    print(f"Excel: {destino.relative_to(RAIZ) if destino.is_relative_to(RAIZ) else destino}")
    print(f"Prospectos nuevos: {len(nuevos)}")
    if repetidos:
        print("Descartados por estar ya en el histórico: " + ", ".join(p.get("empresa", "") for p in repetidos))


if __name__ == "__main__":
    main()
